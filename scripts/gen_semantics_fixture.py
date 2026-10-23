#!/usr/bin/env python3
"""Generate tests/assets/semantics.{xlsx,ods}.

The workbook holds one formula per row of the `Calc` sheet over fixed data in
`Data`. openpyxl writes the formulas without values; LibreOffice then
recalculates and stores them (XLSX and ODS), so the cached results are an
independent oracle for the engine (see tests/oracle.rs).

Also generates tests/assets/loan.{xlsx,ods}, the example model used by the
CLI tests and the README.

Usage: scripts/gen_semantics_fixture.py   (needs openpyxl and soffice)
"""
import os
import subprocess
import tempfile
from datetime import date

import openpyxl

OUT = os.path.join(os.path.dirname(__file__), "..", "tests", "assets")

FORMULAS = [
    # aggregates
    "SUM(Data!A1:A8)", "AVERAGE(Data!A1:A8)", "MIN(Data!A1:A8)", "MAX(Data!A1:A8)",
    "COUNT(Data!A1:A8)", "COUNTA(Data!B1:B9)", "MEDIAN(Data!A1:A8)",
    "LARGE(Data!A1:A8,2)", "SMALL(Data!A1:A8,3)", "PRODUCT(Data!A1:A3)",
    "SUMIF(Data!B1:B8,\"apple\",Data!A1:A8)", "COUNTIF(Data!B1:B8,\"a*\")",
    "COUNTIF(Data!A1:A8,\">4\")", "COUNTIF(Data!A1:A8,\"<>5\")",
    "SUMIFS(Data!A1:A8,Data!B1:B8,\"apple\",Data!A1:A8,\">=5\")",
    "AVERAGEIF(Data!B1:B8,\"banana\",Data!A1:A8)",
    "COUNTIFS(Data!B1:B8,\"apple\",Data!A1:A8,\"<6\")",
    "_xlfn.MAXIFS(Data!A1:A8,Data!B1:B8,\"apple\")",
    "_xlfn.MINIFS(Data!A1:A8,Data!B1:B8,\"apple\")",
    "SUMPRODUCT(Data!A1:A8,Data!A1:A8)",
    "SUMPRODUCT((Data!B1:B8=\"apple\")*(Data!A1:A8))",
    "COUNTBLANK(Data!B1:B9)", "SUM(Data!A:A)", "SUMSQ(Data!A1:A3)",
    "STDEV(Data!A1:A8)", "VAR(Data!A1:A8)",
    # math
    "ROUND(Data!A1/3,2)", "ROUND(2.5,0)", "ROUND(-2.5,0)", "ROUND(1.005,2)", "ROUND(1234.567,-2)",
    "ROUNDUP(2.341,2)", "ROUNDDOWN(-2.349,2)", "TRUNC(8.97)", "TRUNC(8.97,1)",
    "INT(-3.2)", "MOD(-7,3)", "MOD(7,-3)", "POWER(2,10)", "SQRT(16)", "ABS(-4)", "SIGN(-3)",
    "EXP(1)", "LN(10)", "LOG(100,10)", "LOG(8,2)", "LOG10(1000)", "CEILING(4.3,0.5)",
    "FLOOR(4.3,0.5)", "PI()", "EVEN(3)", "ODD(2)", "EVEN(-3)", "FACT(5)", "DEGREES(PI())",
    "SIN(PI()/2)", "ATAN2(1,1)",
    # logic and info
    "IF(Data!A1>3,\"big\",\"small\")", "IF(Data!A2>3,\"big\",\"small\")", "IFERROR(1/0,\"err\")",
    "_xlfn.IFNA(NA(),\"na\")", "AND(TRUE,1,Data!A1>1)", "OR(FALSE,0,Data!A1<1)", "NOT(Data!A1>1)",
    "_xlfn.XOR(TRUE,TRUE)", "CHOOSE(2,\"a\",\"b\",\"c\")", "_xlfn.IFS(Data!A1<1,\"x\",Data!A1>1,\"y\")",
    "_xlfn.SWITCH(Data!A1,5,\"five\",3,\"three\",\"other\")",
    "ISNUMBER(Data!A1)", "ISTEXT(Data!B1)", "ISBLANK(Data!B9)", "ISERROR(1/0)", "ISNA(NA())",
    "ISEVEN(4)", "ISODD(4)", "ISLOGICAL(TRUE)", "N(TRUE)",
    # text
    "LEN(Data!B2)", "UPPER(Data!B2)", "LOWER(\"ABC\")", "PROPER(\"hELLO wORLD\")", "TRIM(\"  a   b  \")",
    "LEFT(Data!B2,3)", "RIGHT(Data!B2,2)", "MID(Data!B2,2,3)", "FIND(\"n\",Data!B2)",
    "SEARCH(\"AN\",Data!B2)", "SEARCH(\"b*a\",Data!B2)", "SUBSTITUTE(\"a-b-c\",\"-\",\"+\")",
    "SUBSTITUTE(\"a-b-c\",\"-\",\"+\",2)", "REPLACE(\"abcdef\",2,3,\"X\")", "REPT(\"ab\",3)",
    "EXACT(\"a\",\"A\")", "CONCATENATE(\"a\",1,\"b\")", "_xlfn.CONCAT(\"a\",Data!A1)",
    "_xlfn.TEXTJOIN(\"-\",TRUE,\"a\",\"\",\"b\")", "VALUE(\"12.5\")", "VALUE(\"50%\")",
    "TEXT(1234.567,\"#,##0.00\")", "TEXT(0.256,\"0.0%\")", "\"[\"&TEXT(5,\"000\")&\"]\"",
    "TEXT(Data!C1,\"yyyy-mm-dd\")", "TEXT(Data!C1,\"dddd\")", "TEXT(Data!C1,\"mmm d, yyyy\")",
    "TEXT(-3.5,\"0.0;(0.0)\")", "CHAR(65)", "CODE(\"A\")", "FIXED(1234.567,1)",
    "\"[\"&(0.1+0.2)&\"]\"", "\"[\"&(1/3)&\"]\"", "\"[\"&100000000000000000000&\"]\"",
    "\"[\"&0.00001&\"]\"", "\"[\"&1234567.891&\"]\"", "\"[\"&-0.5&\"]\"",
    # lookup
    "VLOOKUP(\"cherry\",Data!B1:C8,2,FALSE)", "VLOOKUP(\"zzz\",Data!B1:C8,2,FALSE)",
    "VLOOKUP(4,Data!A1:C8,3,TRUE)", "HLOOKUP(\"y\",Data!E1:G2,2,FALSE)",
    "MATCH(\"cherry\",Data!B1:B8,0)", "MATCH(\"ch*\",Data!B1:B8,0)", "MATCH(6,Data!H1:H5,1)",
    "MATCH(6,Data!I1:I5,-1)", "INDEX(Data!A1:A8,3)", "INDEX(Data!A1:C8,2,1)",
    "SUM(INDEX(Data!A1:C8,0,1))", "SUM(INDEX(Data!A1:C8,2,0))",
    "_xlfn.XLOOKUP(\"cherry\",Data!B1:B8,Data!A1:A8)", "_xlfn.XLOOKUP(\"q\",Data!B1:B8,Data!A1:A8,\"none\")",
    "ROW(Data!A5)", "COLUMN(Data!C1)", "ROWS(Data!A1:C8)", "COLUMNS(Data!A1:C8)",
    "ADDRESS(2,3)", "ADDRESS(2,3,4)", "OFFSET(Data!A1,2,0)", "SUM(OFFSET(Data!A1,0,0,3,1))",
    "INDIRECT(\"Data!A3\")", "SUM(INDIRECT(\"Data!A1:A3\"))", "INDIRECT(\"Data!B\"&2)",
    # dates
    "DATE(2024,1,31)", "DATE(2024,14,1)", "DATE(2024,3,0)", "DATE(1900,3,1)", "DATE(2024,2,29)+1",
    "YEAR(Data!C1)", "MONTH(Data!C1)", "DAY(Data!C1)", "WEEKDAY(Data!C1)", "WEEKDAY(Data!C1,2)",
    "EDATE(DATE(2024,1,31),1)", "EDATE(DATE(2024,3,31),-1)", "EOMONTH(Data!C1,0)", "EOMONTH(Data!C1,2)",
    "_xlfn.DAYS(Data!C2,Data!C1)", "DATEDIF(Data!C1,Data!C2,\"M\")", "DATEDIF(Data!C1,Data!C2,\"D\")",
    "DATEDIF(DATE(2020,1,15),DATE(2024,1,14),\"Y\")", "HOUR(0.75)", "MINUTE(0.7512)", "SECOND(0.7512)",
    "TIME(13,30,0)", "DATEVALUE(\"2024-03-01\")",
    # finance
    "PMT(0.01,12,-1000)", "PMT(0.05/12,360,200000)", "FV(0.05,10,-100)", "PV(0.05,10,-100)",
    "NPV(0.1,100,200,300)", "IRR(Data!J1:J5)",
    # operators and coercion
    "-2^2", "2^-1", "2^3^2", "10%*3", "\"3\"+1", "TRUE+1", "\"a\"&\"b\"", "1=\"1\"", "\"a\"<\"B\"",
    "\"abc\"=\"ABC\"", "1<\"a\"", "(1+2)*3", "7/2", "-(-3)",
    "Data!A1:A3*2", "SUM(Data!A1:A3*2)", "SUM(Data!A1:A3>3)",
    # errors
    "1/0", "SQRT(-1)", "NA()", "\"a\"+1", "0^0", "LN(0)", "VALUE(\"abc\")", "MATCH(99,Data!A1:A8,0)",
    "INDEX(Data!A1:A8,99)", "SUMPRODUCT(LEN(Data!B1:B3))", "SUMPRODUCT(--(Data!A1:A8>4))", "Data!A1+Data!B1", "#REF!+1", "ISERROR(Data!A1/0)",
]


# Formulas returning dates/times are compared through TEXT(): LibreOffice
# stores such results as typed date cells whose serial number the sheetrs ODS
# reader derives from the calendar date (without the 1900 leap-year quirk,
# so it is one less than the spreadsheet serial from 1900-03-01 on).
def as_text(f):
    fmt = "0.0000" if f.startswith("TIME(") else "0"
    return f'"["&TEXT({f},"{fmt}")&"]"'


FORMULAS = [as_text(f) if f.startswith(("DATE(", "EDATE(", "EOMONTH(", "DATEVALUE(", "TIME(")) else f for f in FORMULAS]


def serial(d):
    return (d - date(1899, 12, 30)).days


def build(path):
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    nums = [5, 3, 8, 1, 9, 2, 7, 4]
    labels = ["apple", "banana", "cherry", "apple", "date", "banana", "apple", "fig"]
    for i, (n, t) in enumerate(zip(nums, labels), start=1):
        data[f"A{i}"] = n
        data[f"B{i}"] = t
    base = date(2024, 1, 15)
    for i in range(8):
        data[f"C{i + 1}"] = serial(date(2024, 1 + i, 15 + i))  # plain numbers: see as_text
    data["E1"], data["F1"], data["G1"] = "x", "y", "z"
    data["E2"], data["F2"], data["G2"] = 10, 20, 30
    for i, v in enumerate([1, 3, 5, 7, 9], start=1):
        data[f"H{i}"] = v
    for i, v in enumerate([9, 7, 5, 3, 1], start=1):
        data[f"I{i}"] = v
    for i, v in enumerate([-100, 30, 40, 50, 60], start=1):
        data[f"J{i}"] = v
    calc = wb.create_sheet("Calc")
    for i, f in enumerate(FORMULAS, start=1):
        calc[f"A{i}"] = "=" + f
        calc[f"B{i}"] = f  # label (as text) for failure messages
        calc[f"B{i}"].data_type = "s"
    wb.save(path)


def main():
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "semantics.xlsx")
        build(src)
        outdir = os.path.join(tmp, "out")
        for fmt in ("xlsx", "ods"):
            subprocess.run(
                ["soffice", "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                check=True, capture_output=True,
            )
            os.replace(os.path.join(outdir, f"semantics.{fmt}"), os.path.join(OUT, f"semantics.{fmt}"))
    print(f"{len(FORMULAS)} formulas")


def build_loan(path):
    """A small, realistic model: loan inputs → payment, totals, schedule."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Loan"
    ws["A1"], ws["B1"] = "Principal", 200000
    ws["A2"], ws["B2"] = "Annual rate", 0.05
    ws["A3"], ws["B3"] = "Years", 30
    ws["A4"], ws["B4"] = "Monthly payment", "=PMT(B2/12,B3*12,-B1)"
    ws["A5"], ws["B5"] = "Total paid", "=B4*B3*12"
    ws["A6"], ws["B6"] = "Total interest", "=B5-B1"
    ws["A7"], ws["B7"] = "Month to inspect", 12
    ws["A8"], ws["B8"] = "Balance at that month", '=INDIRECT("D"&(10+B7))'
    ws["A9"], ws["B9"], ws["C9"], ws["D9"] = "Schedule", "Payment", "Interest", "Balance"
    ws["A10"], ws["D10"] = 0, "=B1"
    for m in range(1, 361):
        r = 10 + m
        ws[f"A{r}"] = m
        ws[f"B{r}"] = "=$B$4"
        ws[f"C{r}"] = f"=D{r - 1}*$B$2/12"
        ws[f"D{r}"] = f"=D{r - 1}+C{r}-B{r}"
    ws["A372"], ws["B372"] = "Unused note", "=SUM(B11:B370)*0+LEN(A1)"
    wb.save(path)


def main_loan():
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, "loan.xlsx")
        build_loan(src)
        outdir = os.path.join(tmp, "out")
        for fmt in ("xlsx", "ods"):
            subprocess.run(
                ["soffice", "--headless", "--convert-to", fmt, "--outdir", outdir, src],
                check=True, capture_output=True,
            )
            os.replace(os.path.join(outdir, f"loan.{fmt}"), os.path.join(OUT, f"loan.{fmt}"))


if __name__ == "__main__":
    main()
    main_loan()
